"""
Goal-4 first pass: validate detected RING COUNT (not yet distances/mm --
that needs per-image scale-bar detection to calibrate against Results2.xlsx's
raw pixel measurements) against the known Age column in Results2.xlsx.

Headless -- reuses the same traced-line + CLAHE + smoothing + valley-finding
logic as 02-06, but runs on the 10 known test images without any GUI, using
the hand-picked root points from test_centerline_batch.py so it can run
unattended.

Run:
    ./.venv/bin/python validate_ring_count.py
"""
import os
import sys

import cv2
import numpy as np
import openpyxl
from scipy.signal import find_peaks

import config
from importlib import import_module
draw = import_module("02_draw_lines")

sys.path.insert(0, os.path.dirname(__file__))
from test_centerline_batch import HAND_PICKED_ROOTS

IMG_DIR = os.path.join(
    config.PROJECT_DIR, "Surf Clam hinge images", "NOAA Surf Clam Survey",
    "Reader comparison selected images",
)
RESULTS_XLSX = os.path.join(IMG_DIR, "Results2.xlsx")


def load_known_ages():
    wb = openpyxl.load_workbook(RESULTS_XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    file_i, age_i = header.index("File"), header.index("Age")
    ages = {}
    for row in rows[1:]:
        if row[file_i] is None:
            continue
        fname = os.path.splitext(row[file_i])[0]
        if fname not in ages:  # keep first occurrence only
            ages[fname] = row[age_i]
    return ages


def detect_ring_count(img, path):
    """path: (N,2) float array of traced centerline points."""
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    clahe = cv2.createCLAHE(config.CLAHE_CLIP, (config.CLAHE_TILE, config.CLAHE_TILE))
    gray_clahe = clahe.apply(gray)

    xs = np.clip(path[:, 0].astype(int), 0, gray.shape[1] - 1)
    ys = np.clip(path[:, 1].astype(int), 0, gray.shape[0] - 1)
    profile = gray_clahe[ys, xs].astype(float)

    from scipy.ndimage import gaussian_filter1d
    smooth = gaussian_filter1d(profile, sigma=config.GRAY_SIGMA)
    rng = smooth.max() - smooth.min()
    valleys, _ = find_peaks(-smooth, distance=config.MIN_DIST, prominence=rng * config.PROMINENCE)
    return len(valleys)


def main():
    ages = load_known_ages()
    print(f"{len(ages)} known ages loaded from Results2.xlsx\n")

    errs = []
    for tag, root_xy in HAND_PICKED_ROOTS.items():
        if tag not in ages:
            print(f"SKIP {tag}: not in Results2.xlsx")
            continue
        img = cv2.imread(os.path.join(IMG_DIR, tag + ".jpg"))
        if img is None:
            print(f"SKIP {tag}: unreadable")
            continue
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        mask = draw.segment(gray)
        root_xy = np.array(root_xy, dtype=float)

        edge = draw.centerline_path(mask)
        edge = draw.despike(edge)
        edge = draw.trim_near(edge, root_xy)
        edge = draw.relax_near_root(edge, root_xy, gray)
        path = draw.smooth(edge, sigma=config.CENTERLINE_SMOOTH_SIGMA)
        path = draw.orient_root_first(path, root_xy)
        path = draw.anchor_to_root(path, root_xy)

        detected = detect_ring_count(img, path)
        known = ages[tag]
        err = detected - known
        errs.append(err)
        print(f"{tag:45s} known={known:3d}  detected={detected:3d}  diff={err:+d}")

    errs = np.array(errs)
    print(f"\nn={len(errs)}  mean diff={errs.mean():+.2f}  mean abs diff={np.abs(errs).mean():.2f}  "
          f"exact matches={np.sum(errs==0)}/{len(errs)}")


if __name__ == "__main__":
    main()
